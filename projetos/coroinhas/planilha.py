"""Leitura da planilha da Pastoral (.xlsx): colunas Função, Nome e Data de Nascimento (em qualquer ordem)."""
from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path

try:
    from python_calamine import CalamineWorkbook
except ImportError:                     # no APK (Android) não há python-calamine: usa o openpyxl
    CalamineWorkbook = None

from . import regras
from .repositorio import ErroCoroinhas


def _data(v) -> str | None:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)) and v > 1000:                 # número serial do Excel
        return (date(1899, 12, 30) + timedelta(days=int(v))).isoformat()
    if isinstance(v, str):
        partes = v.strip().split("/")
        if len(partes) == 3 and all(p.isdigit() for p in partes):
            d, m, a = map(int, partes)
            try:
                return date(a, m, d).isoformat()
            except ValueError:
                return None
    return None


def _tipo(v) -> str:
    t = regras.normalizar(v)
    if t.startswith("cerimon"):
        return "cerimoniario"
    if t.startswith("coord"):
        return "coordenacao"
    return "coroinha"


def _linhas(caminho: Path) -> list[list]:
    if CalamineWorkbook is not None:
        return CalamineWorkbook.from_path(str(caminho)).get_sheet_by_index(0).to_python()
    from openpyxl import load_workbook
    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        return [["" if c is None else c for c in linha] for linha in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


def ler_planilha(caminho: Path) -> list[dict]:
    try:
        linhas = _linhas(caminho)
    except Exception:
        raise ErroCoroinhas("Não consegui abrir a planilha. Confira se é um arquivo .xlsx válido.") from None
    i_cab = next((i for i, l in enumerate(linhas) if any(regras.normalizar(c) == "nome" for c in l)), None)
    if i_cab is None:
        raise ErroCoroinhas('Não achei a coluna "Nome" na planilha.')
    cab = [regras.normalizar(c) for c in linhas[i_cab]]
    c_nome = cab.index("nome")
    c_tipo = next((i for i, c in enumerate(cab) if c.startswith(("funcao", "tipo"))), None)
    c_nasc = next((i for i, c in enumerate(cab) if "nascimento" in c), None)

    membros = []
    for linha in linhas[i_cab + 1:]:
        nome = str(linha[c_nome] or "").strip() if c_nome < len(linha) else ""
        if not nome:
            continue
        membros.append({
            "nome": regras.nome_proprio(nome),
            "tipo": _tipo(linha[c_tipo]) if c_tipo is not None else "coroinha",
            "nascimento": _data(linha[c_nasc]) if c_nasc is not None else None,
            "ausente": False,
        })
    if not membros:
        raise ErroCoroinhas("A planilha não tem nenhum nome preenchido.")
    return membros
